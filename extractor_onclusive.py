import io
import re
import html
import unicodedata
from collections import Counter
from urllib.parse import urlparse, urlunparse

import pandas as pd
import streamlit as st
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

try:
    import requests
except Exception:  # requests es opcional para resolución remota de usuarios
    requests = None


REDES_PERMITIDAS = {"X", "FACEBOOK", "INSTAGRAM", "TIKTOK"}


def quitar_acentos(texto):
    if texto is None or (isinstance(texto, float) and pd.isna(texto)):
        return ""
    return "".join(
        c
        for c in unicodedata.normalize("NFD", str(texto))
        if unicodedata.category(c) != "Mn"
    ).lower()


def obtener_campo(row, candidatos):
    columnas = {quitar_acentos(str(c).strip()): c for c in row.index}
    for candidato in candidatos:
        original = columnas.get(quitar_acentos(candidato))
        if original is None:
            continue
        val = row[original]
        if val is None or pd.isna(val):
            continue
        texto = str(val).strip()
        if texto and texto.lower() not in {"nan", "none", "null"}:
            return texto
    return ""


def cargar_onclusive(file):
    raw = file.read()
    file.seek(0)
    try:
        hojas = pd.read_excel(io.BytesIO(raw), sheet_name=None)
        if hojas:
            return pd.concat(hojas.values(), ignore_index=True)
    except Exception:
        pass

    for encoding in ("utf-8-sig", "utf-8", "latin1"):
        try:
            return pd.read_csv(io.BytesIO(raw), encoding=encoding, on_bad_lines="skip")
        except Exception:
            continue
    raise ValueError("No fue posible leer el archivo como Excel o CSV.")


def normalizar_red(valor, link=""):
    t = quitar_acentos(valor)
    u = (link or "").lower()

    if "twitter" in t or re.search(r"(^|\W)x($|\W)", t) or "x.com/" in u or "twitter.com/" in u:
        return "X"
    if "facebook" in t or "fb" == t.strip() or "facebook.com/" in u or "fb.watch/" in u:
        return "FACEBOOK"
    if "instagram" in t or "instagram.com/" in u:
        return "INSTAGRAM"
    if "tiktok" in t or "tiktok.com/" in u or "vm.tiktok.com/" in u:
        return "TIKTOK"
    return ""


def detectar_red(row):
    link = obtener_link_principal(row)
    for campo in [
        "Media type", "Media Type", "Tipo de Medio", "Tipo de medio",
        "Social Network", "Social network", "Red Social", "Fuente", "Canal",
        "Source type", "Platform", "Plataforma"
    ]:
        val = obtener_campo(row, [campo])
        red = normalizar_red(val, link)
        if red:
            return red
    return normalizar_red("", link)


def obtener_link_principal(row):
    candidatos = [
        "Link URL Medio", "Link de Nota", "URL", "Url", "Enlace", "Link",
        "Permalink", "Post URL", "Post Url", "Original URL"
    ]
    urls = []
    for c in candidatos:
        val = obtener_campo(row, [c])
        if val.startswith("http") and val not in urls:
            urls.append(val)
    if not urls:
        for val in row.values:
            if val is None or pd.isna(val):
                continue
            s = str(val).strip()
            if s.startswith("http") and s not in urls:
                urls.append(s)
    if not urls:
        return ""

    externos = [u for u in urls if "hanakua.mx" not in u.lower()]
    return (externos or urls)[0]


def canonicalizar_link(url):
    if not url:
        return ""
    url = html.unescape(str(url).strip())
    try:
        p = urlparse(url)
        host = p.netloc.lower().replace("www.", "")
        path = p.path
        # En X/Facebook /photo/1 o /video/1 puede ser solo una variante del mismo post.
        # En TikTok /video/<id> es parte esencial del permalink y debe conservarse.
        if host in {"x.com", "twitter.com", "mobile.twitter.com", "facebook.com", "m.facebook.com"}:
            path = re.sub(r"/(photo|video)/\d+/?$", "", path, flags=re.I)
        path = re.sub(r"/+", "/", path).rstrip("/")
        if host == "twitter.com":
            host = "x.com"
        return urlunparse((p.scheme or "https", host, path, "", "", ""))
    except Exception:
        return re.sub(r"[?#].*$", "", url).rstrip("/")


def limpiar_texto_publicacion(texto):
    if not texto:
        return ""
    t = str(texto)
    # URLs dentro del contenido: se conserva únicamente el enlace principal por separado.
    t = re.sub(r"https?://\S+", "", t)
    # Sufijos de variante que a veces quedan sueltos al exportar contenido.
    t = re.sub(r"(?<!\w)/(?:photo|video)/\d+\b", "", t, flags=re.I)
    # Emojis escritos como códigos :purple_heart:, :warning:, etc.
    t = re.sub(r":[A-Za-z0-9_+\-|]+:", "", t)
    # IDs numéricos de Facebook como @840385099165296
    t = re.sub(r"(?<!\w)@\d{7,}\b", "", t)
    t = re.sub(r"[ \t]+", " ", t)
    t = re.sub(r"\s*\n\s*", " ", t)
    return t.strip(" -\t\n")


def texto_de_fila(row):
    titulo = obtener_campo(row, ["Titulo", "Título", "Title", "Encabezado", "Tema"])
    contenido = obtener_campo(
        row,
        ["Contenido", "Detail", "Summary", "Síntesis", "Sintesis", "Nota", "Text", "Texto"]
    )
    if contenido and titulo and quitar_acentos(titulo) not in quitar_acentos(contenido):
        bruto = f"{titulo}. {contenido}"
    else:
        bruto = contenido or titulo
    return limpiar_texto_publicacion(bruto)


def normalizar_texto_dup(texto):
    t = quitar_acentos(texto)
    t = re.sub(r"\s+", " ", t)
    t = re.sub(r"[^a-z0-9#@ ]", "", t)
    return t.strip()


def es_rt(row, texto, red):
    if red != "X":
        return False
    tipo = " ".join(
        obtener_campo(row, [c])
        for c in ["Tipo de Nota", "Post type", "Type", "Publication type", "Interaction type"]
    )
    combinado = quitar_acentos(f"{tipo} {texto}").strip()
    return bool(
        re.match(r"^rt\s+@", combinado)
        or " retweet" in f" {combinado}"
        or " retuit" in f" {combinado}"
        or combinado.startswith("retweet")
        or combinado.startswith("retuit")
    )


def parsear_fecha(valor):
    if valor is None or pd.isna(valor):
        return pd.NaT
    return pd.to_datetime(valor, dayfirst=True, errors="coerce")


def obtener_autor(row):
    return obtener_campo(
        row,
        ["Author name", "Autor", "Author", "Nombre del Medio", "Media name", "Fuente", "Page name", "Account name", "Canal"]
    )


def obtener_handle_exportado(row):
    return obtener_campo(
        row,
        ["Author handle (@username)", "Handle", "Username", "Screen Name", "Account", "Perfil", "User name"]
    ).lstrip("@")


def usuario_desde_url(red, url):
    if not url:
        return ""
    try:
        p = urlparse(url)
        host = p.netloc.lower().replace("www.", "")
        partes = [x for x in p.path.split("/") if x]
    except Exception:
        return ""

    if red == "X" and host in {"x.com", "twitter.com", "mobile.twitter.com"} and partes:
        candidato = partes[0]
        if candidato.lower() not in {"i", "intent", "search", "home", "share"}:
            return candidato.lstrip("@")

    if red == "TIKTOK":
        for parte in partes:
            if parte.startswith("@") and len(parte) > 1:
                return parte[1:]

    if red == "INSTAGRAM" and partes:
        reservado = {"p", "reel", "reels", "tv", "stories", "explore", "accounts", "share"}
        if partes[0].lower() not in reservado:
            return partes[0].lstrip("@")

    if red == "FACEBOOK" and partes:
        reservado = {"share", "reel", "watch", "photo", "groups", "events", "story.php", "permalink.php", "profile.php"}
        candidato = partes[0]
        if candidato.lower() not in reservado and not candidato.isdigit():
            return candidato.lstrip("@")
    return ""


@st.cache_data(ttl=86400, show_spinner=False)
def resolver_usuario_publico(url, red):
    """Mejor esfuerzo. No usa login ni intenta evadir bloqueos."""
    if not url or requests is None or red not in {"INSTAGRAM", "TIKTOK"}:
        return ""
    try:
        resp = requests.get(
            url,
            timeout=5,
            allow_redirects=True,
            headers={
                "User-Agent": (
                    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                    "AppleWebKit/537.36 Chrome/124 Safari/537.36"
                )
            },
        )
        if not resp.ok:
            return ""
        final_url = resp.url
        local = usuario_desde_url(red, final_url)
        if local:
            return local

        cuerpo = resp.text[:500000]
        metas = " ".join(re.findall(r'<meta[^>]+(?:content|name|property)=["\'][^"\']+["\'][^>]*>', cuerpo, flags=re.I))
        texto = html.unescape(re.sub(r"<[^>]+>", " ", metas))

        if red == "INSTAGRAM":
            patrones = [
                r"@([A-Za-z0-9._]{2,30})",
                r"([A-Za-z0-9._]{2,30})\s+on\s+Instagram",
            ]
        else:
            patrones = [r"@([A-Za-z0-9._-]{2,40})"]

        for patron in patrones:
            m = re.search(patron, texto, flags=re.I)
            if m:
                return m.group(1)
    except Exception:
        return ""
    return ""


def recortar_por_tema(texto, tema, max_chars=520):
    texto = (texto or "").strip()
    if len(texto) <= max_chars:
        return texto

    tokens = [
        x for x in re.findall(r"[a-z0-9áéíóúñü]+", tema.lower())
        if len(x) >= 4
    ]
    texto_low = texto.lower()
    posiciones = [texto_low.find(t) for t in tokens if texto_low.find(t) >= 0]
    centro = min(posiciones) if posiciones else 0

    inicio = max(0, centro - 120)
    fin = min(len(texto), inicio + max_chars)
    frag = texto[inicio:fin].strip()
    if inicio > 0:
        frag = "…" + frag
    if fin < len(texto):
        frag += "…"
    return frag


def etiqueta_lista(red, autor, handle):
    if red in {"X", "INSTAGRAM", "TIKTOK"}:
        return f"@{handle}" if handle else (autor or red.title())
    return autor or (f"@{handle}" if handle else "Facebook")


def fuente_desglose(red, autor, handle):
    if red == "X":
        if autor and handle:
            return f"{autor} @{handle}"
        return f"@{handle}" if handle else autor or "X"
    if red == "FACEBOOK":
        return autor or (f"@{handle}" if handle else "Facebook")
    if red in {"INSTAGRAM", "TIKTOK"}:
        if autor and handle:
            return f"{autor} @{handle}"
        return f"@{handle}" if handle else autor or red.title()
    return autor or handle or red


def procesar_onclusive(df, tema, resolver_remoto=False):
    registros = []
    stats = Counter()
    vistos_links = set()
    vistos_texto_fuente = set()

    for idx, row in df.iterrows():
        link_raw = obtener_link_principal(row)
        red = detectar_red(row)
        if red not in REDES_PERMITIDAS:
            stats["otras_redes"] += 1
            continue

        texto = texto_de_fila(row)
        if not texto:
            stats["sin_texto"] += 1
            continue
        if es_rt(row, texto, red):
            stats["rt"] += 1
            continue

        link = canonicalizar_link(link_raw)
        autor = obtener_autor(row)
        handle = obtener_handle_exportado(row) or usuario_desde_url(red, link_raw)

        if not handle and resolver_remoto and red in {"INSTAGRAM", "TIKTOK"}:
            handle = resolver_usuario_publico(link_raw, red)

        # Si el "autor" de Facebook es un ID numérico, no lo mostramos.
        if red == "FACEBOOK" and re.fullmatch(r"@?\d{7,}", autor or ""):
            autor = "Facebook"

        fuente_key = quitar_acentos(handle or autor or red)
        texto_key = normalizar_texto_dup(texto)

        if link and link in vistos_links:
            stats["dup_link"] += 1
            continue
        if texto_key and (fuente_key, texto_key) in vistos_texto_fuente:
            stats["dup_texto"] += 1
            continue

        if link:
            vistos_links.add(link)
        if texto_key:
            vistos_texto_fuente.add((fuente_key, texto_key))

        fecha_raw = obtener_campo(
            row,
            ["Publish date", "Fecha", "Date", "Fecha de publicación", "Fecha de publicacion", "Published", "Publication date"]
        )
        fecha = parsear_fecha(fecha_raw)
        texto = recortar_por_tema(texto, tema)

        registros.append({
            "orden": idx,
            "red": red,
            "autor": autor,
            "handle": handle,
            "texto": texto,
            "link": link or link_raw,
            "fecha": fecha,
        })

    # Fechas válidas primero y, dentro de cada fecha, respeta el orden original.
    registros.sort(
        key=lambda r: (
            pd.Timestamp.max if pd.isna(r["fecha"]) else r["fecha"].normalize(),
            r["orden"],
        )
    )
    return registros, stats


def agregar_hipervinculo(parrafo, texto, url):
    part = parrafo.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(color)
    rPr.append(u)
    new_run.append(rPr)

    text = OxmlElement("w:t")
    text.text = texto
    new_run.append(text)
    hyperlink.append(new_run)
    parrafo._p.append(hyperlink)


def crear_word(registros, tema):
    doc = Document()
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(10)

    p = doc.add_paragraph()
    r = p.add_run(f"REDES SOCIALES ({len(registros)})")
    r.bold = True

    etiquetas = [etiqueta_lista(r["red"], r["autor"], r["handle"]) for r in registros]
    conteos = Counter(etiquetas)
    vistos = set()
    for et in etiquetas:
        if et in vistos:
            continue
        vistos.add(et)
        p = doc.add_paragraph()
        p.add_run(f"{et} ({conteos[et]})" if conteos[et] > 1 else et)

    p = doc.add_paragraph()
    r = p.add_run("DESGLOSE")
    r.bold = True

    fecha_actual = object()
    for item in registros:
        fecha_key = None if pd.isna(item["fecha"]) else item["fecha"].strftime("%d.%m.%y")
        if fecha_key != fecha_actual:
            fecha_actual = fecha_key
            if fecha_key:
                p = doc.add_paragraph()
                rr = p.add_run(fecha_key)
                rr.bold = True

        # Un solo párrafo por mención. add_break() = salto manual tipo Shift+Enter.
        p = doc.add_paragraph()
        fuente = fuente_desglose(item["red"], item["autor"], item["handle"])
        rr = p.add_run(fuente)
        rr.bold = True
        rr.add_break()
        p.add_run(item["texto"])
        p.add_run().add_break()
        if item["link"]:
            agregar_hipervinculo(p, item["link"], item["link"])

    out = io.BytesIO()
    doc.save(out)
    out.seek(0)
    return out


def crear_txt(registros):
    lineas = [f"REDES SOCIALES ({len(registros)})"]
    etiquetas = [etiqueta_lista(r["red"], r["autor"], r["handle"]) for r in registros]
    conteos = Counter(etiquetas)
    vistos = set()
    for et in etiquetas:
        if et not in vistos:
            vistos.add(et)
            lineas.append(f"{et} ({conteos[et]})" if conteos[et] > 1 else et)

    lineas += ["", "DESGLOSE"]
    fecha_actual = object()
    for item in registros:
        fecha_key = None if pd.isna(item["fecha"]) else item["fecha"].strftime("%d.%m.%y")
        if fecha_key != fecha_actual:
            fecha_actual = fecha_key
            if fecha_key:
                lineas += ["", fecha_key]
        lineas += [
            fuente_desglose(item["red"], item["autor"], item["handle"]),
            item["texto"],
            item["link"],
            "",
        ]
    return "\n".join(lineas).encode("utf-8")


def crear_html(registros):
    etiquetas = [etiqueta_lista(r["red"], r["autor"], r["handle"]) for r in registros]
    conteos = Counter(etiquetas)
    vistos = set()
    partes = [
        "<!doctype html><html><head><meta charset='utf-8'>",
        "<style>body{font-family:Arial,sans-serif;font-size:14px;line-height:1.35} .m{margin-bottom:14px}</style>",
        "</head><body>",
        f"<p><strong>REDES SOCIALES ({len(registros)})</strong></p>",
    ]
    for et in etiquetas:
        if et in vistos:
            continue
        vistos.add(et)
        txt = f"{et} ({conteos[et]})" if conteos[et] > 1 else et
        partes.append(f"<div>{html.escape(txt)}</div>")

    partes.append("<p><strong>DESGLOSE</strong></p>")
    fecha_actual = object()
    for item in registros:
        fecha_key = None if pd.isna(item["fecha"]) else item["fecha"].strftime("%d.%m.%y")
        if fecha_key != fecha_actual:
            fecha_actual = fecha_key
            if fecha_key:
                partes.append(f"<p><strong>{fecha_key}</strong></p>")
        fuente = html.escape(fuente_desglose(item["red"], item["autor"], item["handle"]))
        texto = html.escape(item["texto"])
        link = html.escape(item["link"] or "")
        link_html = f'<a href="{link}">{link}</a>' if link else ""
        partes.append(f"<div class='m'><strong>{fuente}</strong><br>{texto}<br>{link_html}</div>")
    partes.append("</body></html>")
    return "".join(partes).encode("utf-8")


def render_extractor_onclusive():
    st.subheader("Extracción de menciones Onclusive por tema")
    st.caption(
        "Extrae Facebook, X, Instagram y TikTok; elimina RT/duplicados y genera Word, HTML y TXT."
    )

    tema = st.text_input(
        "Nombre del tema",
        placeholder="Ej. Seguridad en Puebla, Cablebús, Gira de trabajo por Zacatlán",
        key="onclusive_tema",
    ).strip()
    archivo = st.file_uploader(
        "Sube el Excel o CSV exportado desde Onclusive",
        type=["xlsx", "xls", "csv"],
        key="onclusive_archivo",
    )
    resolver = st.checkbox(
        "Intentar recuperar usuarios faltantes de Instagram/TikTok desde publicaciones públicas",
        value=False,
        help=(
            "TikTok suele poder resolverse desde el enlace. En Instagram no siempre es posible; "
            "el intento remoto depende de que la publicación sea pública y accesible sin iniciar sesión."
        ),
        key="onclusive_resolver_usuarios",
    )

    if archivo and tema and st.button("Extraer menciones", type="primary", key="onclusive_extraer"):
        with st.spinner("Depurando publicaciones y preparando entregables..."):
            try:
                df = cargar_onclusive(archivo)
                registros, stats = procesar_onclusive(df, tema, resolver_remoto=resolver)
                if not registros:
                    st.warning("No se encontraron menciones válidas en Facebook, X, Instagram o TikTok.")
                    return

                total_por_red = Counter(r["red"] for r in registros)
                st.success(f"Se extrajeron {len(registros)} menciones válidas.")
                st.write(
                    " · ".join(
                        f"{red.title() if red != 'X' else 'X'}: {total_por_red.get(red, 0)}"
                        for red in ["X", "FACEBOOK", "INSTAGRAM", "TIKTOK"]
                    )
                )

                omitidas = sum(stats.values())
                if omitidas:
                    st.caption(
                        "Omitidas/depuradas: "
                        f"otras redes {stats['otras_redes']}, RT {stats['rt']}, "
                        f"duplicados por link {stats['dup_link']}, "
                        f"duplicados por texto y fuente {stats['dup_texto']}, "
                        f"sin texto {stats['sin_texto']}."
                    )

                faltan_handle = sum(
                    1 for r in registros if r["red"] in {"X", "INSTAGRAM", "TIKTOK"} and not r["handle"]
                )
                if faltan_handle:
                    st.info(
                        f"{faltan_handle} publicación(es) no permitieron identificar un @usuario con los datos disponibles. "
                        "Se conservó el nombre del autor/página cuando estaba disponible."
                    )

                base = re.sub(r"[^A-Za-z0-9ÁÉÍÓÚáéíóúÑñ_-]+", "_", tema).strip("_") or "tema"
                word = crear_word(registros, tema)
                html_bytes = crear_html(registros)
                txt_bytes = crear_txt(registros)

                st.download_button(
                    "📥 Descargar Word",
                    data=word,
                    file_name=f"Extraccion_{base}.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
                st.download_button(
                    "📥 Descargar HTML",
                    data=html_bytes,
                    file_name=f"Extraccion_{base}.html",
                    mime="text/html",
                )
                st.download_button(
                    "📥 Descargar TXT",
                    data=txt_bytes,
                    file_name=f"Extraccion_{base}.txt",
                    mime="text/plain",
                )
            except Exception as exc:
                st.error(f"Error al procesar el archivo: {exc}")
